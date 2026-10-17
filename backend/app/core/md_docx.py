# -*- coding: utf-8 -*-
"""
Markdown → Word(docx) 导出器（通用能力）
- 轻量解析：标题(#/##/###)、列表(-)、表格(| |)、加粗(**)、斜体(*)、分隔线(---)
- 输出 .docx，中文字体默认微软雅黑
- 用于报告/文档的 Markdown 内容 → Word 交付文件
"""
from __future__ import annotations

import re
from pathlib import Path

from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml.ns import qn
from docx.shared import Pt, RGBColor

# 加粗/斜体正则（**bold** / *italic*）
_BOLD_RE = re.compile(r"\*\*(.+?)\*\*")
_ITALIC_RE = re.compile(r"\*(.+?)\*")
# 列表行
_LIST_RE = re.compile(r"^[-*]\s+(.*)")
# 表格分隔行
_TABLE_SEP_RE = re.compile(r"^\|[\s:|-]+\|$")
_TABLE_ROW_RE = re.compile(r"^\|(.+)\|$")


def _inline_runs(text: str) -> list[tuple[str, bool, bool]]:
    """把行内文本拆成 [(文本, 是否加粗, 是否斜体), ...]"""
    runs: list[tuple[str, bool, bool]] = []
    # 先按加粗分割，再在非加粗片段里找斜体
    parts = _BOLD_RE.split(text)
    for i, part in enumerate(parts):
        if not part:
            continue
        if i % 2 == 1:
            runs.append((part, True, False))
        else:
            italic_parts = _ITALIC_RE.split(part)
            for j, ip in enumerate(italic_parts):
                if not ip:
                    continue
                if j % 2 == 1:
                    runs.append((ip, False, True))
                else:
                    runs.append((ip, False, False))
    return runs


def _set_run(run, bold: bool = False, italic: bool = False, size: int | None = None,
             color: RGBColor | None = None):
    run.bold = bold
    run.italic = italic
    if size:
        run.font.size = Pt(size)
    if color:
        run.font.color.rgb = color
    # 中文字体
    run.font.name = "微软雅黑"
    run._element.rPr.rFonts.set(qn("w:eastAsia"), "微软雅黑")


def _add_paragraph(doc: Document, text: str, style: str | None = None):
    p = doc.add_paragraph(style=style) if style else doc.add_paragraph()
    for t, b, i in _inline_runs(text):
        run = p.add_run(t)
        _set_run(run, bold=b, italic=i, size=11)
    return p


def _add_heading(doc: Document, text: str, level: int):
    p = doc.add_heading(level=min(level, 4))
    for t, b, i in _inline_runs(text):
        run = p.add_run(t)
        _set_run(run, bold=True, size=max(16 - level, 12),
                 color=RGBColor(0x1F, 0x3B, 0x63))
    return p


def _parse_table_row(line: str) -> list[str]:
    inner = line.strip().strip("|")
    return [c.strip() for c in inner.split("|")]


def md_to_docx(md_text: str, output_path: str | Path) -> Path:
    """
    将 Markdown 文本导出为 docx 文件。

    Args:
        md_text: Markdown 内容（# 标题、## 章节、- 列表、| 表格）
        output_path: 输出 .docx 路径

    Returns:
        输出文件路径
    """
    output_path = Path(output_path)
    output_path.parent.mkdir(parents=True, exist_ok=True)

    doc = Document()
    # 默认样式
    style = doc.styles["Normal"]
    style.font.name = "微软雅黑"
    style.font.size = Pt(11)
    style.element.rPr.rFonts.set(qn("w:eastAsia"), "微软雅黑")

    lines = md_text.splitlines()
    i = 0
    table_buffer: list[str] = []  # 收集表格行

    def flush_table():
        nonlocal table_buffer
        if not table_buffer:
            return
        header = _parse_table_row(table_buffer[0])
        rows = [_parse_table_row(r) for r in table_buffer[2:]]  # 跳过分隔行
        n_cols = len(header)
        table = doc.add_table(rows=1 + len(rows), cols=n_cols)
        table.style = "Light Grid Accent 1"
        for ci, h in enumerate(header):
            cell = table.rows[0].cells[ci]
            cell.text = ""
            run = cell.paragraphs[0].add_run(h)
            _set_run(run, bold=True, size=10)
        for ri, row in enumerate(rows, 1):
            for ci in range(n_cols):
                val = row[ci] if ci < len(row) else ""
                cell = table.rows[ri].cells[ci]
                cell.text = ""
                run = cell.paragraphs[0].add_run(val)
                _set_run(run, size=10)
        doc.add_paragraph()  # 表格后空行
        table_buffer = []

    for line in lines:
        line = line.rstrip()

        # 表格行：以 | 开头
        if line.lstrip().startswith("|"):
            if table_buffer and _TABLE_SEP_RE.match(line.strip()):
                table_buffer.append(line.strip())
                continue
            table_buffer.append(line.strip())
            continue
        else:
            flush_table()

        line = line.strip()
        if not line:
            continue
        if line.startswith("---"):
            continue  # 分隔线
        m = re.match(r"^(#{1,6})\s+(.*)", line)
        if m:
            _add_heading(doc, m.group(2).strip(), len(m.group(1)))
            continue
        m = _LIST_RE.match(line)
        if m:
            _add_paragraph(doc, m.group(1), style="List Bullet")
            continue
        # 普通段落
        _add_paragraph(doc, line)

    flush_table()
    doc.save(output_path)
    return output_path
